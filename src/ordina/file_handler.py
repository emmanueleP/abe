import os
from PyQt5.QtCore import Qt
from PIL import Image
import fitz  # PyMuPDF
from docx import Document
import openpyxl
from .utils import create_stamp, get_output_path
from .settings import ordina_settings as settings
import io
from .history_dialog import add_to_history

def handle_file(file_path):
    """Gestisce il file in base al suo tipo"""
    try:
        # Ottieni il percorso di output
        output_path = get_output_path(file_path)
        print(f"Output path: {output_path}")  # Debug
        
        # Ottieni il nome del file
        base_name = os.path.basename(output_path)
        print(f"Base name: {base_name}")  # Debug
        
        # Verifica che il nome del file contenga '__'
        if '__' not in base_name:
            raise Exception("Nome file non valido: manca il separatore '__'")
            
        # Split del nome file e verifica lunghezza
        parts = base_name.split('__')
        if len(parts) != 2:
            raise Exception(f"Nome file non valido: formato errato ({base_name})")
            
        # Ottieni il numero di protocollo
        protocol_parts = parts[1].split('.')
        if len(protocol_parts) < 1:
            raise Exception(f"Nome file non valido: manca il numero di protocollo ({parts[1]})")
            
        protocol_number = protocol_parts[0]
        print(f"Protocol number: {protocol_number}")  # Debug
        
        # Crea il timbro
        stamp = create_stamp(protocol_number)
        if stamp is None:
            raise Exception("Errore nella creazione del timbro")
            
        # Gestisci il file in base all'estensione
        ext = os.path.splitext(file_path)[1].lower()
        print(f"File extension: {ext}")  # Debug
        
        if ext == '.pdf':
            handle_pdf(file_path, output_path, stamp)
        elif ext == '.docx':
            handle_docx(file_path, output_path, stamp)
        elif ext == '.xlsx':
            handle_xlsx(file_path, output_path, stamp)
        elif ext in ['.png', '.jpg', '.jpeg']:
            handle_image(file_path, output_path, stamp)
        else:
            raise Exception(f"Formato file non supportato: {ext}")
        
        # Aggiungi alla cronologia
        add_to_history(protocol_number, output_path)
        
        return output_path
        
    except Exception as e:
        print(f"DEBUG - Errore dettagliato: {str(e)}")  # Debug
        raise Exception(f"Errore durante la protocollazione: {str(e)}")

def handle_pdf(input_path, output_path, stamp):
    """Gestisce file PDF"""
    try:
        # Apri il PDF
        pdf = fitz.open(input_path)
        
        # Converti il timbro PIL in bytes
        stamp_bytes = io.BytesIO()
        stamp.save(stamp_bytes, format='PNG')
        stamp_bytes = stamp_bytes.getvalue()
        
        # Inserisci il timbro nella prima pagina
        first_page = pdf[0]
        
        # Calcola la posizione del timbro
        stamp_position = settings.current_settings.get("stamp_position", "top-right")
        rect = first_page.rect
        stamp_width = stamp.width / 2  # Converti da pixel a punti PDF
        stamp_height = stamp.height / 2
        
        if stamp_position == "top-right":
            x = rect.width - stamp_width - 20
            y = 20
        elif stamp_position == "top-left":
            x = 20
            y = 20
        elif stamp_position == "bottom-right":
            x = rect.width - stamp_width - 20
            y = rect.height - stamp_height - 20
        else:  # bottom-left
            x = 20
            y = rect.height - stamp_height - 20
        
        # Inserisci il timbro
        first_page.insert_image((x, y, x + stamp_width, y + stamp_height), stream=stamp_bytes)
        
        # Salva il PDF
        pdf.save(output_path)
        pdf.close()
        
    except Exception as e:
        raise Exception(f"Errore nella gestione del PDF: {str(e)}")

def handle_docx(input_path, output_path, stamp):
    """Gestisce file Word"""
    doc = Document(input_path)
    
    # Salva il timbro come immagine temporanea
    temp_stamp = "temp_stamp.png"
    stamp.save(temp_stamp, "PNG")
    
    # Aggiungi il timbro al documento
    doc.add_picture(temp_stamp)
    
    # Salva il documento
    doc.save(output_path)
    
    # Rimuovi il file temporaneo
    os.remove(temp_stamp)

def handle_xlsx(input_path, output_path, stamp):
    """Gestisce file Excel"""
    wb = openpyxl.load_workbook(input_path)
    ws = wb.active
    
    # Salva il timbro come immagine temporanea
    temp_stamp = "temp_stamp.png"
    stamp.save(temp_stamp, "PNG")
    
    # Aggiungi il timbro al foglio
    img = openpyxl.drawing.image.Image(temp_stamp)
    ws.add_image(img, 'A1')
    
    # Salva il file
    wb.save(output_path)
    
    # Rimuovi il file temporaneo
    os.remove(temp_stamp)

def handle_image(input_path, output_path, stamp):
    """Gestisce file immagine"""
    with Image.open(input_path) as img:
        # Converti in RGBA se necessario
        if img.mode != 'RGBA':
            img = img.convert('RGBA')
        
        # Calcola la posizione del timbro (in basso a destra)
        x = img.width - stamp.width - 50
        y = img.height - stamp.height - 50
        
        # Incolla il timbro
        img.paste(stamp, (x, y), stamp)
        
        # Salva l'immagine
        img.save(output_path)